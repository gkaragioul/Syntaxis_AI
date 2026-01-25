from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any
import re

class ExcelWriter:
    def __init__(self):
        self.wb = Workbook()
    
    def write_grid(self, grid: List[List[str]], sheet_name: str = "Sheet1", 
                   header_depth: int = 0, normalize_numbers: bool = True,
                   keep_currency: bool = False, merges: List[Dict] = None):
        
        if "Sheet" in self.wb.sheetnames:
            ws = self.wb.active
            ws.title = sheet_name
        else:
            ws = self.wb.create_sheet(sheet_name)
        
        for row_idx, row_data in enumerate(grid, start=1):
            for col_idx, cell_value in enumerate(row_data, start=1):
                processed_value = cell_value
                
                if normalize_numbers and row_idx > header_depth:
                    processed_value = self.normalize_number(cell_value, keep_currency)
                
                ws.cell(row=row_idx, column=col_idx, value=processed_value)
                
                if row_idx <= header_depth:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        
        if merges:
            for merge in merges:
                start_row = merge.get("start_row", 1)
                end_row = merge.get("end_row", 1)
                start_col = merge.get("start_col", 1)
                end_col = merge.get("end_col", 1)
                
                ws.merge_cells(
                    start_row=start_row, start_column=start_col,
                    end_row=end_row, end_column=end_col
                )
        
        if header_depth > 0:
            ws.freeze_panes = ws.cell(row=header_depth + 1, column=1)
        
        for col_idx in range(1, len(grid[0]) + 1 if grid else 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 15
    
    def normalize_number(self, value: str, keep_currency: bool = False) -> Any:
        if not value:
            return value
        
        original = value
        
        is_negative = False
        if value.startswith("(") and value.endswith(")"):
            is_negative = True
            value = value[1:-1]
        
        if not keep_currency:
            value = re.sub(r'[$€£¥₹]', '', value)
        
        value = value.replace(',', '').strip()
        
        try:
            num = float(value)
            if is_negative:
                num = -num
            return num
        except ValueError:
            return original
    
    def save(self, filepath: str):
        self.wb.save(filepath)
    
    def close(self):
        self.wb.close()
